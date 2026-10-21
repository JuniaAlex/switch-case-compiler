#!/usr/bin/env python3
"""Build docs/PROJECT_REPORT.docx from docs/PROJECT_REPORT.md."""

import re
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
OUTPUT = ROOT / "docs" / "PROJECT_REPORT.docx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def border_paragraph(paragraph, color="0F766E", size="18", side="left"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    edge = OxmlElement("w:" + side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "6")
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)
    p_pr.append(p_bdr)


def table_borders(table, color="CBD5E1"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + edge_name)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)
    tbl_pr.append(borders)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    tr_pr.append(cant)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90):
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement("w:" + name)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = BODY_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    _field(paragraph, "PAGE")
    run = paragraph.add_run(" of ")
    run.font.name = BODY_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    _field(paragraph, "NUMPAGES")


def _field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " %s " % instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "2-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field to refresh the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = SLATE
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Heading 2", 16, NAVY, 0, 10),
        ("Heading 3", 12.5, BLUE, 12, 5),
        ("Heading 4", 11, SLATE, 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 2":
            style.paragraph_format.page_break_before = True

    toc_heading = doc.styles["TOC Heading"]
    toc_heading.font.name = HEAD_FONT
    toc_heading.font.size = Pt(18)
    toc_heading.font.bold = True
    toc_heading.font.color.rgb = NAVY
    toc_heading.paragraph_format.space_after = Pt(10)


def paragraph_border_bottom(paragraph, color="0F766E", size="18"):
    border_paragraph(paragraph, color=color, size=size, side="bottom")


def add_inline(paragraph, node, bold=False, italic=False, code=False, size=None):
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = TEAL
        elif size:
            run.font.size = Pt(size)
        return

    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name in ("strong", "b"):
        for child in node.children:
            add_inline(paragraph, child, True, italic, code, size)
    elif name in ("em", "i"):
        for child in node.children:
            add_inline(paragraph, child, bold, True, code, size)
    elif name == "code":
        run = paragraph.add_run(node.get_text())
        run.font.name = "Consolas"
        run.font.size = Pt((size or 10.5) - 1)
        run.font.color.rgb = TEAL
        run.bold = bold
    elif name == "br":
        paragraph.add_run().add_break()
    elif name in ("a", "span", "sup", "sub"):
        for child in node.children:
            add_inline(paragraph, child, bold, italic, code, size)
    else:
        for child in node.children:
            add_inline(paragraph, child, bold, italic, code, size)


def add_code_block(doc, text):
    text = text.strip("\n")
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shade_paragraph(paragraph, "F1F5F9")
    border_paragraph(paragraph, color="0F766E", size="16", side="left")
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.6)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)


def add_table(doc, table_tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    column_count = max(len(row.find_all(["th", "td"])) for row in rows)
    table = doc.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    set_cell_margins(table)
    keep_table = len(rows) <= 16

    for row_index, row_tag in enumerate(rows):
        cells_tags = row_tag.find_all(["th", "td"])
        row = table.add_row()
        cant_split(row)
        if row_index == 0:
            repeat_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if column_index < len(cells_tags):
                add_inline(paragraph, cells_tags[column_index], size=8.8)
            if keep_table and row_index < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
            if row_index == 0:
                shade(cell, "0F172A")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(8.8)
            elif row_index % 2 == 0:
                shade(cell, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def image_path(src):
    return (SOURCE.parent / src).resolve()


def add_image(doc, tag):
    src = tag.get("src", "")
    if not src:
        return
    path = image_path(src)
    if not path.exists():
        print("warning: missing image", path)
        return
    with Image.open(path) as image:
        width, height = image.size
    max_width = 6.15
    max_height = 7.0
    width_in = max_width
    height_in = width_in * height / width
    if height_in > max_height:
        height_in = max_height
        width_in = height_in * width / height

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width_in), height=Inches(height_in))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    run = caption.add_run(tag.get("alt", "Figure"))
    run.italic = True
    run.font.size = Pt(8.8)
    run.font.color.rgb = MUTED


def add_list(doc, tag, ordered=False, level=0):
    for index, item in enumerate(tag.find_all("li", recursive=False), start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.32 + 0.24 * level)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        prefix = ("%d. " % index) if ordered else "\u2022 "
        run = paragraph.add_run(prefix)
        run.bold = ordered
        nested = []
        for child in item.children:
            if isinstance(child, Tag) and child.name in ("ul", "ol"):
                nested.append(child)
            else:
                add_inline(paragraph, child, size=10.2)
        for nested_tag in nested:
            add_list(doc, nested_tag, ordered=nested_tag.name == "ol", level=level + 1)


def add_body(doc, soup):
    for child in soup.children:
        if isinstance(child, NavigableString):
            if child.strip():
                paragraph = doc.add_paragraph(str(child).strip())
            continue
        if not isinstance(child, Tag):
            continue
        name = child.name.lower()
        if name in ("h2", "h3", "h4"):
            level = int(name[1])
            doc.add_heading(child.get_text(" ", strip=True), level=level)
        elif name == "p":
            images = child.find_all("img")
            if images and not child.get_text(strip=True):
                for image in images:
                    add_image(doc, image)
                continue
            paragraph = doc.add_paragraph()
            add_inline(paragraph, child)
        elif name == "ul":
            add_list(doc, child, ordered=False)
        elif name == "ol":
            add_list(doc, child, ordered=True)
        elif name == "pre":
            add_code_block(doc, child.get_text())
        elif name == "table":
            add_table(doc, child)
        elif name == "blockquote":
            for p in child.find_all("p"):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.3)
                run = paragraph.add_run(p.get_text(" ", strip=True))
                run.italic = True
                run.font.color.rgb = MUTED
        elif name == "hr":
            paragraph = doc.add_paragraph()
            paragraph_border_bottom(paragraph, color="CBD5E1", size="6")
        elif name == "div":
            add_body(doc, child)


def cover_metadata(soup):
    h1 = soup.find("h1")
    subtitle_tag = h1.find_next_sibling("p") if h1 else None
    table_tag = h1.find_next_sibling("table") if h1 else None
    title = h1.get_text(" ", strip=True) if h1 else "SQL Mini Compiler"
    subtitle = subtitle_tag.get_text(" ", strip=True) if subtitle_tag else ""
    metadata = []
    if table_tag:
        for row in table_tag.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if len(cells) >= 2:
                key = cells[0].get_text(" ", strip=True)
                value = cells[1].get_text(" ", strip=True)
                if key.lower() != "field":
                    metadata.append((key, value))
    for tag in (h1, subtitle_tag, table_tag):
        if tag is not None:
            tag.decompose()
    return title, subtitle, metadata


def build_cover(doc, title, subtitle, metadata):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "SQL Mini Compiler  |  Compiler Design Project"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_p.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
    border_paragraph(header_p, color="CBD5E1", size="4", side="bottom")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)

    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(18)
    label.paragraph_format.space_after = Pt(4)
    run = label.add_run("COMPILER DESIGN PROJECT")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = TEAL
    run.font.small_caps = True

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.line_spacing = 1.0
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title_p.add_run(title)
    run.bold = True
    run.font.size = Pt(34)
    run.font.name = HEAD_FONT
    run.font.color.rgb = NAVY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(6)
    subtitle_p.paragraph_format.space_after = Pt(8)
    run = subtitle_p.add_run(subtitle)
    run.font.size = Pt(15)
    run.font.color.rgb = MUTED
    run.italic = True

    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(12)
    paragraph_border_bottom(line, color="0F766E", size="22")

    snippet = "SELECT name, cgpa\nFROM students\nWHERE age >= 18\nORDER BY cgpa DESC;"
    code_p = doc.add_paragraph()
    code_p.paragraph_format.left_indent = Inches(0.1)
    code_p.paragraph_format.right_indent = Inches(0.1)
    code_p.paragraph_format.space_before = Pt(2)
    code_p.paragraph_format.space_after = Pt(18)
    code_p.paragraph_format.line_spacing = 1.0
    code_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shade_paragraph(code_p, "F1F5F9")
    border_paragraph(code_p, color="0F766E", size="14", side="left")
    for index, line_text in enumerate(snippet.splitlines()):
        if index:
            code_p.add_run().add_break()
        run = code_p.add_run(line_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEAL

    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.autofit = False
    for key, value in metadata:
        row = meta_table.add_row()
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(4.55)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(key)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = TEAL
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        r1 = p1.add_run(value)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = SLATE

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(12)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Validation only. No query execution and no database connection.")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = MUTED

    doc.add_page_break()


def build():
    markdown_text = SOURCE.read_text(encoding="utf-8")
    html = markdown.markdown(
        markdown_text,
        extensions=["tables", "fenced_code", "sane_lists", "extra"],
    )
    soup = BeautifulSoup(html, "html.parser")
    title, subtitle, metadata = cover_metadata(soup)

    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = title
    doc.core_properties.subject = "Compiler Design project report"
    doc.core_properties.author = "SQL Mini Compiler Team"
    doc.core_properties.keywords = "SQL, compiler, lexer, parser, Flex, Bison, semantic analysis"

    build_cover(doc, title, subtitle, metadata)

    contents = doc.add_paragraph(style="TOC Heading")
    contents.add_run("Contents")
    toc_paragraph = doc.add_paragraph()
    add_toc(toc_paragraph)

    add_body(doc, soup)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("wrote", OUTPUT)
    print("pages are paginated when the PDF is exported; the TOC field is updated by Word")


if __name__ == "__main__":
    build()
